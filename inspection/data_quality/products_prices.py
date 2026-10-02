"""
inspection/data_quality/products_prices.py

Metziah data quality inspection report for products, store products,
and prices.

Generates an Excel workbook containing database overview, per-chain
price/item analysis, and chain-unique item analysis.
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from db import get_connection


REPORT_DIR = Path("inspection/reports")
SHEETS_DIR = Path("inspection/sheets")


def fetch_one(cursor, query):
    cursor.execute(query)
    row = cursor.fetchone()
    return row[0] if row else 0


def fetch_all(cursor, query):
    cursor.execute(query)
    return cursor.fetchall()


def get_overview(cursor):
    return {
        "Chains": fetch_one(cursor, "SELECT COUNT(*) FROM chains"),
        "Stores": fetch_one(cursor, "SELECT COUNT(*) FROM stores"),
        "Products": fetch_one(cursor, "SELECT COUNT(*) FROM products"),
        "Store Products": fetch_one(
            cursor,
            "SELECT COUNT(*) FROM store_products",
        ),
        "Price Rows": fetch_one(cursor, "SELECT COUNT(*) FROM prices"),
        "Distinct Price Item Codes": fetch_one(
            cursor,
            "SELECT COUNT(DISTINCT item_code) FROM prices",
        ),
    }


def get_by_chain(cursor):
    query = """
        SELECT
            c.name_en_normalized AS chain,
            COUNT(DISTINCT p.store_id) AS stores,
            COUNT(*) AS price_rows,
            COUNT(DISTINCT p.item_code) AS distinct_item_codes,
            COUNT(DISTINCT p.item_code) FILTER (
                WHERE EXISTS (
                    SELECT 1
                    FROM products pr
                    WHERE pr.item_code = p.item_code
                )
            ) AS item_codes_in_products,
            COUNT(DISTINCT p.item_code) FILTER (
                WHERE NOT EXISTS (
                    SELECT 1
                    FROM products pr
                    WHERE pr.item_code = p.item_code
                )
            ) AS item_codes_not_in_products
        FROM prices p
        JOIN chains c
            ON c.chain_id = p.chain_id
        GROUP BY
            c.chain_id,
            c.name_en_normalized
        ORDER BY price_rows DESC
    """

    return fetch_all(cursor, query)


def get_chain_unique_items(cursor):
    query = """
        WITH chain_items AS (
            SELECT
                chain_id,
                item_code
            FROM prices
            GROUP BY
                chain_id,
                item_code
        ),
        item_chain_counts AS (
            SELECT
                item_code,
                COUNT(*) AS chain_count
            FROM chain_items
            GROUP BY item_code
        ),
        unique_chain_items AS (
            SELECT
                ci.chain_id,
                ci.item_code
            FROM chain_items ci
            JOIN item_chain_counts icc
                ON icc.item_code = ci.item_code
            WHERE icc.chain_count = 1
        )
        SELECT
            c.name_en_normalized AS chain,
            COUNT(*) AS unique_item_codes,
            COUNT(*) FILTER (
                WHERE EXISTS (
                    SELECT 1
                    FROM products p
                    WHERE p.item_code = uci.item_code
                )
            ) AS unique_in_products,
            COUNT(*) FILTER (
                WHERE NOT EXISTS (
                    SELECT 1
                    FROM products p
                    WHERE p.item_code = uci.item_code
                )
            ) AS unique_not_in_products
        FROM unique_chain_items uci
        JOIN chains c
            ON c.chain_id = uci.chain_id
        GROUP BY
            c.chain_id,
            c.name_en_normalized
        ORDER BY unique_item_codes DESC
    """

    return fetch_all(cursor, query)


def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            "solid",
            fgColor="D9EAF7",
        )

    for column_cells in ws.columns:
        max_length = 0

        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))

        width = min(max(max_length + 2, 10), 40)

        ws.column_dimensions[
            get_column_letter(column_cells[0].column)
        ].width = width


def write_table(ws, headers, rows):
    ws.append(headers)

    for row in rows:
        ws.append(list(row))

    style_sheet(ws)


def create_workbook(data):
    wb = Workbook()

    default = wb.active
    wb.remove(default)

    # ---------------------------------------------------------
    # Overview
    # ---------------------------------------------------------

    ws = wb.create_sheet("Overview")

    ws.append(["Metric", "Count"])

    for metric, value in data["overview"].items():
        ws.append([metric, value])

    style_sheet(ws)

    # ---------------------------------------------------------
    # By Chain
    # ---------------------------------------------------------

    write_table(
        wb.create_sheet("By Chain"),
        [
            "Chain",
            "Stores",
            "Price Rows",
            "Distinct Item Codes",
            "In Products",
            "Not In Products",
        ],
        data["by_chain"],
    )

    # ---------------------------------------------------------
    # Chain-Unique Items
    # ---------------------------------------------------------

    write_table(
        wb.create_sheet("Chain-Unique Items"),
        [
            "Chain",
            "Unique Item Codes",
            "Unique In Products",
            "Unique Not In Products",
        ],
        data["chain_unique_items"],
    )

    return wb


def main():
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    SHEETS_DIR.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y-%m-%d")

    xlsx_path = (
        SHEETS_DIR
        / f"products_prices_{timestamp}.xlsx"
    )

    with get_connection() as conn:
        with conn.cursor() as cursor:
            data = {
                "overview": get_overview(cursor),
                "by_chain": get_by_chain(cursor),
                "chain_unique_items": get_chain_unique_items(cursor),
            }

    workbook = create_workbook(data)
    workbook.save(xlsx_path)

    print(f"Excel report: {xlsx_path}")


if __name__ == "__main__":
    main()